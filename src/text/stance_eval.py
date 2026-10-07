"""Benchmark de postura hawkish/dovish · Bloque 2, fase 2.

- Datos de referencia: WCB (frases del BCE etiquetadas por expertos, de las actas de
  política monetaria) y muestra propia de ruedas de prensa, etiquetada a ciegas.
- Predicciones de los candidatos, con caché en disco.
- Métricas, intervalos de confianza por bootstrap y regla de decisión.
"""
import math
import re
import time
import unicodedata
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import spearmanr
from sklearn.metrics import accuracy_score, f1_score, precision_recall_fscore_support

from src.text.models.stance_backends import COLUMNAS, liberar_memoria
from src.text.stance import etiquetar

# Esquema de etiquetas del WCB, que adoptamos también en nuestro etiquetado
ETIQUETAS_ANOTACION = ["hawkish", "dovish", "neutral", "irrelevant"]

# Clases de evaluación: en la app, "irrelevant" se muestra como neutral
ETIQUETAS_EVALUACION = ["hawkish", "neutral", "dovish"]

WCB_DATASET = "gtfintechlab/european_central_bank"
WCB_SEMILLA = "5768"  # partición train/val/test que usamos (la primera de las tres del WCB)

# Iniciales admitidas al etiquetar, por si se escribe a mano en lugar de usar el desplegable
INICIALES = {"h": "hawkish", "d": "dovish", "n": "neutral", "i": "irrelevant"}

# Guía de etiquetado: definiciones del WCB adaptadas al BCE (se copia en el Excel)
GUIA_ETIQUETADO = [
    ("hawkish",
     "Apunta a una política más restrictiva, o a mantenerla restrictiva, para contener la inflación. "
     "Señales: subidas de tipos o sesgo alcista; reducir el balance (fin de compras netas o de reinversiones); "
     "inflación alta o al alza; riesgos al alza para la inflación; salarios, empleo o crecimiento fuertes que "
     "presionan los precios; subida de los precios de la energía; depreciación del euro."),
    ("dovish",
     "Apunta a una política más acomodaticia, o a mantenerla acomodaticia. "
     "Señales: bajadas de tipos o sesgo bajista; compras de activos, TLTRO u otras medidas de liquidez; "
     "inflación baja o a la baja; riesgos a la baja para el crecimiento o la inflación; actividad débil o "
     "desempleo al alza; caída de los precios de la energía; apreciación del euro."),
    ("neutral",
     "Habla de economía o de política monetaria sin una dirección clara: señales en ambos sentidos o "
     "equilibradas, compromisos genéricos (\"we will follow a data-dependent and meeting-by-meeting approach\", "
     "\"we stand ready to adjust all of our instruments\") o descripciones sin implicación para los precios."),
    ("irrelevant",
     "No aporta información sobre política monetaria ni sobre la economía: saludos, agradecimientos, "
     "logística, procedimiento, anécdotas."),
]

REGLAS_ETIQUETADO = [
    "Etiqueta solo lo que dice la frase: no uses el contexto ni lo que sabes que pasó después.",
    "Los datos económicos cuentan por la presión que ejercen sobre los precios: "
    "\"inflation rose to 3.0 per cent\" → hawkish; \"growth slowed\" → dovish.",
    "Si hay señales en ambos sentidos, elige la dominante; si están equilibradas, neutral.",
    "Mantener los tipos, sin más, es neutral; si se añade una orientación (\"for as long as necessary\", "
    "\"at their present or lower levels\"), etiqueta según esa orientación.",
    "Si dudas, elige la etiqueta más probable y escribe \"duda\" en la columna de comentario.",
]


# ---------------------------------------------------------------------------
# WCB: frases del BCE etiquetadas por expertos
# ---------------------------------------------------------------------------
def normalizar(texto: pd.Series) -> pd.Series:
    """Minúsculas y solo letras y números, para comparar frases de fuentes distintas."""
    return (texto.str.lower()
                 .str.replace(r"['’]", "", regex=True)          # "don't" -> "dont"
                 .str.replace(r"[^a-z0-9]+", " ", regex=True)    # resto de signos -> espacio
                 .str.strip())


def cargar_wcb(semilla: str = WCB_SEMILLA) -> pd.DataFrame:
    """Frases del BCE del WCB con su partición (train/val/test) y su etiqueta de postura."""
    from datasets import load_dataset  # solo se necesita en el benchmark

    ds = load_dataset(WCB_DATASET, semilla)
    partes = []
    for split in ds:
        df = ds[split].to_pandas()
        df["split"] = split
        partes.append(df)
    wcb = pd.concat(partes, ignore_index=True)
    wcb = wcb.rename(columns={"sentences": "text", "stance_label": "label"})
    wcb["label"] = wcb["label"].str.strip().str.lower()
    return wcb[["split", "year", "text", "label"]]


def solapamiento(wcb: pd.DataFrame, corpus: pd.DataFrame) -> set[str]:
    """Frases (normalizadas) que aparecen a la vez en el WCB y en nuestro corpus."""
    return set(normalizar(wcb["text"])) & set(normalizar(corpus["text"]))


# ---------------------------------------------------------------------------
# Muestra propia de ruedas de prensa
# ---------------------------------------------------------------------------
def muestra_en_dominio(corpus: pd.DataFrame, por_rueda: dict[str, int] | None = None,
                       min_palabras: int = 6, excluir: set[str] = frozenset(),
                       semilla: int = 42) -> pd.DataFrame:
    """Muestra estratificada de frases del panel para etiquetar.

    De cada rueda de prensa se toman, al azar, frases de la declaración y de las
    respuestas (presidenta o vicepresidente) con un mínimo de palabras. El orden final
    es aleatorio, así que cualquier tramo inicial del fichero es también una muestra
    aleatoria (útil si no da tiempo a etiquetarlo entero).
    """
    por_rueda = por_rueda or {"statement": 1, "answer": 3}

    panel = corpus[corpus["speaker"].isin(["presidenta", "vicepresidente"])
                   & corpus["role"].isin(list(por_rueda))]
    panel = panel[panel["text"].str.split().str.len() >= min_palabras]
    panel = panel[~normalizar(panel["text"]).isin(excluir)]   # fuera las frases que estén en el WCB

    partes = [panel[panel["role"] == rol].groupby("date", group_keys=False).sample(n=n, random_state=semilla)
              for rol, n in por_rueda.items()]
    muestra = pd.concat(partes).sample(frac=1, random_state=semilla).reset_index(drop=True)
    muestra.insert(0, "id", range(1, len(muestra) + 1))
    return muestra[["id", "date", "sentence_id", "section", "role", "speaker", "text"]]


def crear_excel_etiquetado(muestra: pd.DataFrame, ruta: Path) -> None:
    """Excel para etiquetar a ciegas: sin fecha ni hablante, con desplegable de etiquetas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor="E8E6DF")
    ajuste = Alignment(wrap_text=True, vertical="top")

    libro = Workbook()

    # Hoja 1: instrucciones
    guia = libro.active
    guia.title = "Instrucciones"
    guia.column_dimensions["A"].width = 14
    guia.column_dimensions["B"].width = 110
    guia.append(["Pregunta", "¿Qué orientación de política monetaria transmite la frase, leída de forma aislada?"])
    guia.append([])
    guia.append(["Etiqueta", "Cuándo usarla"])
    for etiqueta, descripcion in GUIA_ETIQUETADO:
        guia.append([etiqueta, descripcion])
    guia.append([])
    guia.append(["Reglas", ""])
    for i, regla in enumerate(REGLAS_ETIQUETADO, start=1):
        guia.append([f"{i}.", regla])
    guia.append([])
    guia.append(["Cómo", "Etiqueta en la hoja \"Etiquetado\" con el desplegable de la columna D. "
                         "El orden es aleatorio: si no terminas, lo etiquetado hasta ese punto sigue valiendo."])
    for fila in guia.iter_rows():
        for celda in fila:
            celda.alignment = ajuste
        if fila[0].value in ("Pregunta", "Etiqueta", "Reglas", "Cómo"):
            fila[0].font = negrita
    for celda in guia[3]:
        celda.font = negrita
        celda.fill = relleno
    guia.sheet_properties.pageSetUpPr.fitToPage = True   # si se imprime, que quepa a lo ancho
    guia.page_setup.fitToWidth, guia.page_setup.fitToHeight = 1, 0

    # Hoja 2: etiquetado
    hoja = libro.create_sheet("Etiquetado")
    hoja.append(["id", "seccion", "frase", "etiqueta", "comentario"])
    for fila in muestra.itertuples():
        seccion = "Declaración" if fila.role == "statement" else "Respuesta"
        hoja.append([fila.id, seccion, fila.text, None, None])
        # Alto aproximado de fila para que la frase se lea entera sin ajustar nada
        hoja.row_dimensions[hoja.max_row].height = 15 * math.ceil(len(fila.text) / 95) + 3

    for columna, ancho in zip("ABCDE", [6, 13, 100, 13, 30]):
        hoja.column_dimensions[columna].width = ancho
    for celda in hoja[1]:
        celda.font = negrita
        celda.fill = relleno
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = ajuste
    hoja.freeze_panes = "A2"
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 0

    # Desplegable con las cuatro etiquetas (impide valores no válidos)
    desplegable = DataValidation(type="list", formula1='"hawkish,dovish,neutral,irrelevant"', allow_blank=True)
    desplegable.error = "Elige hawkish, dovish, neutral o irrelevant"
    desplegable.errorTitle = "Etiqueta no válida"
    hoja.add_data_validation(desplegable)
    desplegable.add(f"D2:D{hoja.max_row}")

    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)


def leer_etiquetas(ruta_excel: Path, ruta_muestra: Path) -> pd.DataFrame:
    """Une las etiquetas del Excel con la muestra (fecha, sección y texto de cada frase)."""
    etiquetado = pd.read_excel(ruta_excel, sheet_name="Etiquetado")
    valor = etiquetado["etiqueta"].astype("string").str.strip().str.lower().replace(INICIALES)

    no_validas = etiquetado.loc[valor.notna() & ~valor.isin(ETIQUETAS_ANOTACION), "id"].tolist()
    if no_validas:
        raise ValueError(f"Etiquetas no válidas en las frases con id {no_validas}")

    etiquetado["label"] = valor
    muestra = pd.read_csv(ruta_muestra)
    return muestra.merge(etiquetado[["id", "label", "comentario"]], on="id", how="left")


# ---------------------------------------------------------------------------
# Predicciones de los candidatos (con caché en disco)
# ---------------------------------------------------------------------------
# Frases de control: casos evidentes que permiten detectar un mapa de etiquetas mal
# configurado antes de evaluar (p. ej., hawkish y dovish intercambiados)
FRASES_CONTROL = pd.DataFrame({
    "id": [1, 2, 3, 4],
    "esperado": ["hawkish", "dovish", "neutral", "irrelevant"],
    "text": [
        "The Governing Council today decided to raise the three key ECB interest rates by 50 basis points.",
        "The Governing Council today decided to lower the three key ECB interest rates by 25 basis points.",
        "We will continue to follow a data-dependent and meeting-by-meeting approach to determining "
        "the appropriate monetary policy stance.",
        "Good afternoon, and welcome to our press conference.",
    ],
})


def _nombre_fichero(nombre: str) -> str:
    """"BGE + regresión logística" -> "bge_regresion_logistica"."""
    sin_tildes = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", sin_tildes.lower()).strip("_")


def predecir_candidato(nombre: str, crear, conjuntos: dict[str, pd.DataFrame], carpeta: Path) -> pd.DataFrame:
    """Predice los conjuntos con un candidato y guarda probabilidades, tiempo y memoria.

    Si el candidato ya se ejecutó (existe su fichero), reutiliza el resultado: así el
    notebook se puede volver a ejecutar sin repetir los modelos lentos. Para repetir un
    candidato, basta con borrar su fichero pred_<nombre>.csv.
    """
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / f"pred_{_nombre_fichero(nombre)}.csv"
    if ruta.exists():
        print(f"{nombre}: reutilizo {ruta.name}")
        return pd.read_csv(ruta)

    modelo = crear()
    # Las frases de control sirven también de calentamiento: la primera llamada a la GPU es lenta
    partes = [modelo.predecir(FRASES_CONTROL["text"].tolist()).assign(conjunto="control",
                                                                    id=FRASES_CONTROL["id"].values)]
    n_frases, inicio = 0, time.perf_counter()
    for conjunto, df in conjuntos.items():
        partes.append(modelo.predecir(df["text"].tolist()).assign(conjunto=conjunto, id=df["id"].values))
        n_frases += len(df)
    segundos = time.perf_counter() - inicio

    pred = pd.concat(partes, ignore_index=True)
    pred.insert(0, "modelo", nombre)
    pred = pred[["modelo", "conjunto", "id"] + COLUMNAS]
    pred.to_csv(ruta, index=False)

    tiempo = pd.DataFrame([{
        "modelo": nombre, "frases": n_frases, "segundos": round(segundos, 1),
        "frases_por_segundo": round(n_frases / segundos, 2), "memoria_gb": round(modelo.memoria_gb(), 2),
        # Solo en los LLM: probabilidad media que el modelo deja en las cuatro etiquetas
        "masa_etiquetas": round(float(np.mean(modelo.masa)), 3) if hasattr(modelo, "masa") else np.nan,
    }])
    ruta_tiempos = carpeta / "tiempos.csv"
    if ruta_tiempos.exists():
        anteriores = pd.read_csv(ruta_tiempos)
        tiempo = pd.concat([anteriores[anteriores["modelo"] != nombre], tiempo], ignore_index=True)
    tiempo.to_csv(ruta_tiempos, index=False)

    print(f"{nombre}: {n_frases} frases en {segundos:.0f} s")
    del modelo
    liberar_memoria()
    return pred


def cargar_resultados(carpeta: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Predicciones de todos los candidatos ejecutados y su tabla de tiempos y memoria."""
    predicciones = pd.concat([pd.read_csv(r) for r in sorted(carpeta.glob("pred_*.csv"))], ignore_index=True)
    return predicciones, pd.read_csv(carpeta / "tiempos.csv")


def comprobar_control(predicciones: pd.DataFrame) -> pd.DataFrame:
    """Etiqueta (de las cuatro) que cada modelo asigna a las frases de control, frente a la esperada."""
    control = predicciones[predicciones["conjunto"] == "control"].copy()
    control["predicha"] = control[COLUMNAS].fillna(-1.0).idxmax(axis=1).str.removeprefix("p_")
    tabla = control.pivot(index="modelo", columns="id", values="predicha")
    tabla.columns = [f"{e} (esperada)" for e in FRASES_CONTROL["esperado"]]
    return tabla


# ---------------------------------------------------------------------------
# Métricas
# ---------------------------------------------------------------------------
# Etiqueta ordenada para la correlación con la puntuación (score = p_hawkish − p_dovish)
ORDEN = {"hawkish": 1, "neutral": 0, "irrelevant": 0, "dovish": -1}


def a_tres_clases(etiquetas: pd.Series) -> pd.Series:
    """En la evaluación, "irrelevant" cuenta como neutral (así se muestra en la app)."""
    return etiquetas.replace({"irrelevant": "neutral"})


def _unir(predicciones: pd.DataFrame, referencia: pd.DataFrame, conjunto: str, modelo: str) -> pd.DataFrame:
    """Predicciones de un modelo en un conjunto, unidas a la etiqueta de referencia (columna gold)."""
    pred = predicciones[(predicciones["modelo"] == modelo) & (predicciones["conjunto"] == conjunto)]
    gold = referencia[["id", "label"]].rename(columns={"label": "gold"})
    return gold.merge(etiquetar(pred.reset_index(drop=True)), on="id")


def metricas(datos: pd.DataFrame) -> dict:
    """Métricas de un modelo a partir de la etiqueta de referencia (gold) y su predicción."""
    gold3 = a_tres_clases(datos["gold"])
    f1_clases = f1_score(gold3, datos["label"], labels=ETIQUETAS_EVALUACION, average=None, zero_division=0)
    pred4 = datos[COLUMNAS].fillna(-1.0).idxmax(axis=1).str.removeprefix("p_")
    with warnings.catch_warnings():   # puntuación constante: la correlación no está definida (NaN)
        warnings.simplefilter("ignore")
        correlacion = spearmanr(datos["gold"].map(ORDEN), datos["score"])[0]
    return {
        "f1_macro": f1_score(gold3, datos["label"], labels=ETIQUETAS_EVALUACION, average="macro",
                             zero_division=0),
        **{f"f1_{e}": v for e, v in zip(ETIQUETAS_EVALUACION, f1_clases)},
        "exactitud": accuracy_score(gold3, datos["label"]),
        "spearman": correlacion,
        # Cuatro clases y ponderado por soporte: la métrica del paper del WCB
        "f1_ponderado_4": f1_score(datos["gold"], pred4, labels=ETIQUETAS_ANOTACION, average="weighted",
                                   zero_division=0),
    }


def _indices(etiquetas: pd.Series) -> np.ndarray:
    return etiquetas.map({c: i for i, c in enumerate(ETIQUETAS_EVALUACION)}).to_numpy()


def _f1_macro_lotes(gold: np.ndarray, pred: np.ndarray) -> np.ndarray:
    """F1 macro de muchas remuestras a la vez (cada fila de las matrices es una remuestra)."""
    f1 = []
    for c in range(len(ETIQUETAS_EVALUACION)):
        tp = ((gold == c) & (pred == c)).sum(axis=-1)
        fp = ((gold != c) & (pred == c)).sum(axis=-1)
        fn = ((gold == c) & (pred != c)).sum(axis=-1)
        denominador = 2 * tp + fp + fn
        f1.append(np.where(denominador > 0, 2 * tp / np.maximum(denominador, 1), 0.0))
    return np.mean(f1, axis=0)


def _remuestras(n_frases: int, n: int, semilla: int) -> np.ndarray:
    """Índices de n remuestras con reemplazo (la misma semilla da las mismas remuestras)."""
    return np.random.default_rng(semilla).integers(0, n_frases, size=(n, n_frases))


def intervalo_f1(datos: pd.DataFrame, n: int = 2000, semilla: int = 0) -> tuple[float, float]:
    """Intervalo al 95 % del F1 macro, remuestreando las frases con reemplazo."""
    gold, pred = _indices(a_tres_clases(datos["gold"])), _indices(datos["label"])
    idx = _remuestras(len(gold), n, semilla)
    return tuple(np.percentile(_f1_macro_lotes(gold[idx], pred[idx]), [2.5, 97.5]))


def evaluar(predicciones: pd.DataFrame, referencia: pd.DataFrame, conjunto: str) -> pd.DataFrame:
    """Tabla de métricas de todos los modelos en un conjunto, ordenada por F1 macro."""
    filas = []
    for modelo in predicciones["modelo"].unique():
        datos = _unir(predicciones, referencia, conjunto, modelo)
        inferior, superior = intervalo_f1(datos)
        filas.append({"modelo": modelo, "n": len(datos), **metricas(datos),
                      "ic95_inf": inferior, "ic95_sup": superior})
    columnas = ["modelo", "n", "f1_macro", "ic95_inf", "ic95_sup", "f1_hawkish", "f1_neutral", "f1_dovish",
                "exactitud", "spearman", "f1_ponderado_4"]
    return pd.DataFrame(filas)[columnas].sort_values("f1_macro", ascending=False).reset_index(drop=True)


def comparar_con_mejor(predicciones: pd.DataFrame, referencia: pd.DataFrame, conjunto: str,
                       mejor: str, n: int = 2000, semilla: int = 0) -> pd.DataFrame:
    """Diferencia de F1 macro entre el mejor modelo y cada uno de los demás, con bootstrap pareado.

    Pareado: las dos predicciones se evalúan sobre las mismas remuestras de frases, de modo
    que la dificultad de las frases, común a ambos modelos, no infla la incertidumbre.
    Empate: el intervalo al 95 % de la diferencia incluye el 0.
    """
    base = _unir(predicciones, referencia, conjunto, mejor)
    gold, pred_mejor = _indices(a_tres_clases(base["gold"])), _indices(base["label"])
    idx = _remuestras(len(gold), n, semilla)
    filas = []
    for modelo in predicciones["modelo"].unique():
        if modelo == mejor:
            continue
        otro = _unir(predicciones, referencia, conjunto, modelo).set_index("id").loc[base["id"]]
        pred_otro = _indices(otro["label"].reset_index(drop=True))
        diferencias = _f1_macro_lotes(gold[idx], pred_mejor[idx]) - _f1_macro_lotes(gold[idx], pred_otro[idx])
        observada = _f1_macro_lotes(gold[None], pred_mejor[None])[0] - _f1_macro_lotes(gold[None], pred_otro[None])[0]
        inferior, superior = np.percentile(diferencias, [2.5, 97.5])
        filas.append({"modelo": modelo, "diferencia": observada, "ic95_inf": inferior, "ic95_sup": superior,
                      "empate": inferior <= 0})
    return pd.DataFrame(filas)


def aplicar_regla(tabla: pd.DataFrame, comparacion: pd.DataFrame, tiempos: pd.DataFrame) -> str:
    """Regla de decisión: entre el mejor y los empatados con él, el más rápido."""
    empatados = [tabla.loc[0, "modelo"]]
    if len(comparacion):
        empatados += comparacion.loc[comparacion["empate"], "modelo"].tolist()
    candidatos = tiempos[tiempos["modelo"].isin(empatados)]
    return candidatos.sort_values("frases_por_segundo", ascending=False).iloc[0]["modelo"]


def metricas_relevancia(predicciones: pd.DataFrame, referencia: pd.DataFrame, conjunto: str) -> pd.DataFrame:
    """Detección de frases irrelevantes, solo en los modelos que contemplan esa clase."""
    filas = []
    for modelo in predicciones["modelo"].unique():
        datos = _unir(predicciones, referencia, conjunto, modelo)
        if datos["p_irrelevant"].isna().all():
            continue
        precision, recall, f1, _ = precision_recall_fscore_support(
            datos["gold"].eq("irrelevant"), ~datos["relevant"], average="binary", zero_division=0)
        filas.append({"modelo": modelo, "precision": precision, "recall": recall, "f1": f1})
    return pd.DataFrame(filas).sort_values("f1", ascending=False).reset_index(drop=True)


def matriz_confusion(predicciones: pd.DataFrame, referencia: pd.DataFrame, conjunto: str,
                     modelo: str) -> pd.DataFrame:
    """Filas: etiqueta de referencia (tres clases); columnas: etiqueta predicha."""
    datos = _unir(predicciones, referencia, conjunto, modelo)
    return pd.crosstab(a_tres_clases(datos["gold"]).rename("referencia"), datos["label"].rename("predicha"))
